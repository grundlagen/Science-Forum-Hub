#!/usr/bin/env python3
"""Fingerprint PDF/DOCX/XLSX bid files for common-origin screening.

High similarity is a LEAD, not evidence of collusion. Always falsify common
solicitation templates, portals, scanners, consultants, disclosed partnerships,
and agency-side recombination.
"""
from __future__ import annotations
import argparse, collections, datetime as dt, hashlib, json, math, re, zipfile
from pathlib import Path

SUPPORTED={'.pdf','.docx','.xlsx','.xlsm'}

def h(b): return hashlib.sha256(b).hexdigest()
def hf(p): return h(Path(p).read_bytes())
def norm(v):
    if v is None:return None
    if isinstance(v,(str,int,float,bool)):return v
    if isinstance(v,(dt.datetime,dt.date,dt.time)):return v.isoformat()
    return str(v)
def jac(a,b):
    a,b=set(a or []),set(b or [])
    return len(a&b)/len(a|b) if a and b else 0.0

def textstyle(t):
    w=[x.lower() for x in re.findall(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ'’\-]{2,}",t or '')]
    sh=[]
    for n in (4,5):
        sh += [' '.join(w[i:i+n]) for i in range(max(0,len(w)-n+1)) if len(' '.join(w[i:i+n]))>=28]
    sh=sorted((h(x.encode())[:16],x) for x in set(sh))[:100]
    return {'words':len(w),'shingles':[x for _,x in sh],
            'punct':{c:(t or '').count(c) for c in [';','—','–','…','“','”','’']}}

def pdffp(p):
    import fitz, io
    from PIL import Image
    d=fitz.open(p); pages=[]; fonts=[]; subsets=[]; dq=[]; text=[]
    for pg in d:
        r=pg.rect; pages.append((round(r.width,2),round(r.height,2),int(pg.rotation)))
        for f in pg.get_fonts(full=True):
            base=str(f[3] if len(f)>3 else ''); name=str(f[4] if len(f)>4 else '')
            fonts.append((base,name)); m=re.match(r'^([A-Z]{6})\+',base)
            if m: subsets.append(m.group(1))
        for im in pg.get_images(full=True):
            try:
                b=d.extract_image(im[0]).get('image')
                with Image.open(io.BytesIO(b)) as z:
                    if z.quantization: dq.append(h(json.dumps(z.quantization,sort_keys=True).encode()))
            except Exception: pass
        try:text.append(pg.get_text('text'))
        except Exception:pass
    return {'kind':'pdf','meta':{k:v for k,v in (d.metadata or {}).items() if v},
            'geometry':pages,'fonts':sorted(set(fonts)),'subsets':sorted(set(subsets)),
            'jpeg_dqt':sorted(set(dq)),'text':textstyle('\n'.join(text))}

def ziptime(z):
    c=collections.Counter(i.date_time for i in z.infolist())
    return [(list(k),v) for k,v in c.most_common(10)]

def docxfp(p):
    from docx import Document
    d=Document(p); cp=d.core_properties
    with zipfile.ZipFile(p) as z:
        names=set(z.namelist()); read=lambda n:z.read(n) if n in names else b''
        settings=(read('word/settings.xml')+read('word/document.xml')).decode('utf8','ignore')
        rels=read('word/_rels/settings.xml.rels').decode('utf8','ignore')
        return {'kind':'docx','core':{'author':norm(cp.author),'last_modified_by':norm(cp.last_modified_by),
                    'created':norm(cp.created),'modified':norm(cp.modified)},
                'styles':h(read('word/styles.xml')) if read('word/styles.xml') else None,
                'theme':h(read('word/theme/theme1.xml')) if read('word/theme/theme1.xml') else None,
                'templates':re.findall(r'attachedTemplate[^>]*Target="([^"]+)"',rels),
                'rsids':sorted(set(re.findall(r'rsid\w*="([0-9A-Fa-f]+)"',settings)))[:150],
                'zip_times':ziptime(z),'text':textstyle('\n'.join(x.text for x in d.paragraphs))}

def xlsxfp(p):
    import openpyxl
    wb=openpyxl.load_workbook(p,data_only=False,keep_links=True); pr=wb.properties; text=[]; formulas=[]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str):
                    (formulas if c.value.startswith('=') else text).append(c.value)
    with zipfile.ZipFile(p) as z:
        names=set(z.namelist()); read=lambda n:z.read(n) if n in names else b''
        xml=read('xl/workbook.xml').decode('utf8','ignore'); styles=read('xl/styles.xml')
        dn=[]
        for attrs,val in re.findall(r'<definedName\b([^>]*)>(.*?)</definedName>',xml,re.S):
            nm=re.search(r'name="([^"]+)"',attrs); hd=re.search(r'hidden="([^"]+)"',attrs)
            dn.append((nm.group(1) if nm else None,hd.group(1) if hd else None,re.sub(r'\s+',' ',val)[:200]))
        return {'kind':'xlsx','core':{'creator':norm(pr.creator),'lastModifiedBy':norm(pr.lastModifiedBy),
                    'created':norm(pr.created),'modified':norm(pr.modified)},
                'styles':h(styles) if styles else None,'defined_names':dn,
                'formulas':collections.Counter(re.sub(r'\d+','#',x) for x in formulas).most_common(100),
                'zip_times':ziptime(z),'text':textstyle('\n'.join(text))}

def fp(p):
    p=Path(p); d={'path':str(p),'name':p.name,'bytes':p.stat().st_size,'sha256':hf(p)}
    try:
        d.update(pdffp(p) if p.suffix.lower()=='.pdf' else docxfp(p) if p.suffix.lower()=='.docx' else xlsxfp(p))
    except Exception as e:d['error']=f'{type(e).__name__}: {e}'
    return d

def prevalence(xs):
    out={}
    for k in ('pdf','docx','xlsx'):
        g=[x for x in xs if x.get('kind')==k and not x.get('error')]; C=collections.defaultdict(collections.Counter)
        for x in g:
            if k=='pdf':
                for f in ('creator','producer','author'):
                    if x['meta'].get(f): C[f][x['meta'][f]]+=1
                for v in set(map(tuple,x['fonts'])):C['fonts'][v]+=1
                for v in set(x['subsets']):C['subsets'][v]+=1
                for v in set(x['jpeg_dqt']):C['jpeg_dqt'][v]+=1
            elif k=='docx':
                for f in ('author','last_modified_by'):
                    if x['core'].get(f):C[f][x['core'][f]]+=1
                if x.get('styles'):C['styles'][x['styles']]+=1
                for v in set(x['rsids']):C['rsids'][v]+=1
            else:
                for f in ('creator','lastModifiedBy'):
                    if x['core'].get(f):C[f][x['core'][f]]+=1
                if x.get('styles'):C['styles'][x['styles']]+=1
                for v in {(a,b) for a,b,_ in x['defined_names']}:C['defined_names'][v]+=1
        out[k]=(len(g),C)
    return out

def rare(P,k,f,v):
    N,C=P[k]
    if N<=1:return 1.0
    df=C[f].get(v,1); return max(.05,math.log((N+1)/(df+.5))/math.log(N+1))
def rare_inter(P,k,f,a,b):
    z=set(a or [])&set(b or [])
    return sum(rare(P,k,f,v) for v in z)/len(z) if z else 0

def score(a,b,P):
    if a['sha256']==b['sha256']:return 1.0,['byte-identical']
    if a.get('kind')!=b.get('kind'):return 0,[]
    k=a['kind']; s=0; why=[]
    if k=='pdf':
        for f,w in [('creator',.12),('producer',.08),('author',.10)]:
            v=a['meta'].get(f)
            if v and v==b['meta'].get(f):r=rare(P,k,f,v);s+=w*r;why.append(f'same {f} ({r:.2f} rare): {v}')
        q=jac(map(tuple,a['fonts']),map(tuple,b['fonts']))
        if q>.65:r=rare_inter(P,k,'fonts',map(tuple,a['fonts']),map(tuple,b['fonts']));s+=.12*q*r;why.append(f'font Jaccard {q:.2f}')
        q=jac(a['subsets'],b['subsets'])
        if q>.25:r=rare_inter(P,k,'subsets',a['subsets'],b['subsets']);s+=.16*q*r;why.append(f'font subset Jaccard {q:.2f}')
        q=jac(a['jpeg_dqt'],b['jpeg_dqt'])
        if q>.2:r=rare_inter(P,k,'jpeg_dqt',a['jpeg_dqt'],b['jpeg_dqt']);s+=.16*q*r;why.append(f'JPEG DQT Jaccard {q:.2f}')
    elif k=='docx':
        for f,w in [('author',.12),('last_modified_by',.14)]:
            v=a['core'].get(f)
            if v and v==b['core'].get(f):r=rare(P,k,f,v);s+=w*r;why.append(f'same {f}: {v}')
        if a.get('styles') and a['styles']==b.get('styles'):s+=.12*rare(P,k,'styles',a['styles']);why.append('same styles.xml')
        q=jac(a['rsids'],b['rsids'])
        if q>.2:s+=.14*q*rare_inter(P,k,'rsids',a['rsids'],b['rsids']);why.append(f'RSID Jaccard {q:.2f}')
    else:
        for f,w in [('creator',.12),('lastModifiedBy',.14)]:
            v=a['core'].get(f)
            if v and v==b['core'].get(f):s+=w*rare(P,k,f,v);why.append(f'same {f}: {v}')
        if a.get('styles') and a['styles']==b.get('styles'):s+=.16*rare(P,k,'styles',a['styles']);why.append('same styles.xml')
        aa={(x,y) for x,y,_ in a['defined_names']};bb={(x,y) for x,y,_ in b['defined_names']};q=jac(aa,bb)
        if q>.2:s+=.18*q*rare_inter(P,k,'defined_names',aa,bb);why.append(f'defined-name Jaccard {q:.2f}')
    q=jac(a.get('text',{}).get('shingles'),b.get('text',{}).get('shingles'))
    if q>.08:s+=min(.18,.18*q);why.append(f'text shingle Jaccard {q:.2f}')
    return min(1,s),why

def collect(ps):
    z=[]
    for x in ps:
        p=Path(x); z += [p] if p.is_file() and p.suffix.lower() in SUPPORTED else ([q for q in p.rglob('*') if q.suffix.lower() in SUPPORTED] if p.is_dir() else [])
    return sorted(set(z))

def main():
    ap=argparse.ArgumentParser();ap.add_argument('paths',nargs='+');ap.add_argument('--out',default='fingerprints.json');ap.add_argument('--pairs',type=float,default=.18)
    ns=ap.parse_args();xs=[fp(p) for p in collect(ns.paths)];P=prevalence(xs);pairs=[]
    for i in range(len(xs)):
        for j in range(i+1,len(xs)):
            s,w=score(xs[i],xs[j],P)
            if s>=ns.pairs:pairs.append({'a':xs[i]['path'],'b':xs[j]['path'],'score':round(s,4),'reasons':w})
    pairs.sort(key=lambda x:-x['score']);Path(ns.out).write_text(json.dumps({'warning':'screening leads only','files':xs,'candidate_pairs':pairs},indent=2,default=str))
    print(f'{len(xs)} files; {len(pairs)} pairs >= {ns.pairs}')
    for x in pairs[:25]:print(x['score'],Path(x['a']).name,'<>',Path(x['b']).name,'; '.join(x['reasons'][:6]))
if __name__=='__main__':main()
