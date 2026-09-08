#!/usr/bin/env python3
"""Generate a small calibration corpus for document-workshop screening."""
from pathlib import Path
import tempfile
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
R=Path(tempfile.mkdtemp(prefix='procurement-calibration-'))
def pdf(p,a,t):
 c=canvas.Canvas(str(p));c.setAuthor(a);c.setCreator('OddWorkshopTool 7.13');c.drawString(72,720,t);c.save()
def doc(p,a,t):
 d=Document();d.core_properties.author=a;d.core_properties.last_modified_by=a;d.add_paragraph(t);d.save(p)
def xls(p,a,t):
 w=Workbook();w.properties.creator=a;w.properties.lastModifiedBy=a;w.active['A1']=t;w.active['B1']='=SUM(1,2)';w.create_named_range('WorkshopHiddenRange',w.active,'$A$1');w.save(p)
pdf(R/'alpha.pdf','rare-author-17','peculiar copper semaphore quotation phrase')
pdf(R/'beta.pdf','rare-author-17','peculiar copper semaphore quotation phrase')
pdf(R/'control.pdf','ordinary-user','ordinary unrelated quotation')
doc(R/'eta.docx','rare-doc-author','singular aldermanic pavement wording repeated here')
doc(R/'theta.docx','rare-doc-author','singular aldermanic pavement wording repeated here')
xls(R/'delta.xlsx','rare-sheet-author','unusual municipal lattice schedule')
xls(R/'epsilon.xlsx','rare-sheet-author','unusual municipal lattice schedule')
print(R)
