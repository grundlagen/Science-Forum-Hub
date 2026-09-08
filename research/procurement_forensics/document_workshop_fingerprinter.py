#!/usr/bin/env python3
"""Document-workshop fingerprinting prototype.

Extracts provenance/structure/style fingerprints from PDF, DOCX and XLSX files
and scores pairwise common-origin signals. A high score is a lead for review,
NOT evidence of collusion: shared solicitation templates, consultants, scanners,
software stacks and agency compilation can all produce legitimate similarity.

Usage:
  python document_workshop_fingerprinter.py FILE_OR_DIR [...] --out report.json
  python document_workshop_fingerprinter.py bids/ --pairs 0.25 --csv pairs.csv
"""
from __future__ import annotations
import argparse, collections, datetime as dt, hashlib, json, math, re, zipfile
from pathlib import Path
from typing import Any

SUPPORTED={'.pdf','.docx','.xlsx','.xlsm'}

def sha256_bytes(b:bytes)->str: return hashlib.sha256(b).hexdigest()
def sha256_file(p:Path)->str:
    h=hashlib.sha256()
    with p.open('rb') as f:
        for c in iter(lambda:f.read(1<<20),b''): h.update(c)
    return h.hexdigest()

def norm(v:Any):
    if v is None: return None
    if isinstance(v,(str,int,float,bool)): return v
    if isinstance(v,(dt.datetime,dt.date,dt.time)): return v.isoformat()
    return str(v)

def zip_times(z:zipfile.ZipFile):
    vals=[f"{i.date_time[0]:04}-{i.date_time[1]:02}-{i.date_time[2]:02}T{i.date_time[3]:02}:{i.date_time[4]:02}:{i.date_time[5]:02}" for i in z.infolist()]
    c=collections.Counter(vals)
    return {'distinct':len(c),'top':c.most_common(8)}

def xml_hash(z,name):
    try:return sha256_bytes(z.read(name))
    except KeyError:return None

def text_style(text:str):
    if not text: return {}
    words=re.findall(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ'’\-]{2,}",text)
    lower=[w.lower() for w in words]
    decimals={
      'decimal_point':len(re.findall(r'\b\d+\.\d+\b',text)),
      'decimal_comma':len(re.findall(r'\b\d+,\d+\b',text)),
      'thousands_comma':len(re.findall(r'\b\d{1,3}(?:,\d{3})+\b',text)),
      'thousands_space':len(re.findall(r'\b\d{1,3}(?:[ \u00a0]\d{3})+\b',text)),
    }
    punct={c:text.count(c) for c in [';','—','–','…','“','”','’']}
    shingles=[]
    for n in (4,5):
        for i in range(max(0,len(lower)-n+1)):
            s=' '.join(lower[i:i+n])
            if len(s)>=28: shingles.append(s)
    sh=sorted((sha256_bytes(s.encode())[:16],s) for s in set(shingles))[:80]
    return {'word_count':len(words),'decimal':decimals,'punct':punct,
            'peculiar_tokens':collections.Counter(w for w in lower if len(w)>=10).most_common(25),
            'shingle_sample':[s for _,s in sh]}

def pdf_fp(p:Path):
    import fitz, io
    from PIL import Image
    d=fitz.open(p)
    meta={k:norm(v) for k,v in (d.metadata or {}).items() if v}
    try: trailer_id=d.xref_get_key(-1,'ID')[1]
    except Exception: trailer_id=None
    pages=[]; fonts=[]; image_dqt=[]; texts=[]
    for page in d:
        r=page.rect; pages.append([round(r.width,3),round(r.height,3),int(page.rotation)])
        try:
            for f in page.get_fonts(full=True):
                base=str(f[3] if len(f)>3 else ''); name=str(f[4] if len(f)>4 else '')
                subset=re.match(r'^([A-Z]{6})\+',base)
                fonts.append({'base':base,'name':name,'subset_prefix':subset.group(1) if subset else None})
        except Exception: pass
        try:
            for im in page.get_images(full=True):
                try:
                    b=d.extract_image(im[0]).get('image')
                    if not b: continue
                    with Image.open(io.BytesIO(b)) as pi:
                        q=pi.quantization
                        if q:
                            qcanon=json.dumps({str(k):list(v) for k,v in sorted(q.items())},sort_keys=True).encode()
                            image_dqt.append(sha256_bytes(qcanon))
                except Exception: pass
        except Exception: pass
        try:texts.append(page.get_text('text'))
        except Exception: pass
    return {'kind':'pdf','metadata':meta,'trailer_id':trailer_id,'page_geometry':pages,
            'font_pairs':sorted(set((x['base'],x['name']) for x in fonts)),
            'font_subset_prefixes':sorted(set(x['subset_prefix'] for x in fonts if x['subset_prefix'])),
            'jpeg_quant_hashes':sorted(collections.Counter(image_dqt).items()),'text_style':text_style('\n'.join(texts))}

def docx_fp(p:Path):
    from docx import Document
    doc=Document(p); cp=doc.core_properties
    props={k:norm(getattr(cp,k,None)) for k in ['author','last_modified_by','created','modified','title','subject','category','revision']}
    txt='\n'.join(x.text for x in doc.paragraphs)
    with zipfile.ZipFile(p) as z:
        names=set(z.namelist()); settings=z.read('word/settings.xml').decode('utf-8','ignore') if 'word/settings.xml' in names else ''
        rels=z.read('word/_rels/settings.xml.rels').decode('utf-8','ignore') if 'word/_rels/settings.xml.rels' in names else ''
        template_targets=re.findall(r'Type="[^"]*attachedTemplate[^"]*"[^>]*Target="([^"]+)"',rels)
        docxml=z.read('word/document.xml').decode('utf-8','ignore') if 'word/document.xml' in names else ''
        rsids=sorted(set(re.findall(r'w:rsid(?:R|RPr|Sect|Del|P)?="([0-9A-Fa-f]+)"',settings+docxml)))
        return {'kind':'docx','core':props,'template_targets':template_targets,'styles_hash':xml_hash(z,'word/styles.xml'),
                'theme_hash':xml_hash(z,'word/theme/theme1.xml'),'numbering_hash':xml_hash(z,'word/numbering.xml'),
                'rsids':rsids[:100],'zip_times':zip_times(z),'text_style':text_style(txt)}

def xlsx_fp(p:Path):
    import openpyxl
    wb=openpyxl.load_workbook(p,read_only=False,data_only=False,keep_links=True); props=wb.properties
    core={k:norm(getattr(props,k,None)) for k in ['creator','lastModifiedBy','created','modified','title','subject','category','revision']}
    formulas=[]; text=[]
    sheets=[{'title':ws.title,'state':ws.sheet_state,'max_row':ws.max_row,'max_col':ws.max_column} for ws in wb.worksheets]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v=c.value
                if isinstance(v,str):
                    if v.startswith('='): formulas.append(re.sub(r'\d+','#',v))
                    else:text.append(v)
    with zipfile.ZipFile(p) as z:
        names=set(z.namelist()); workbook=z.read('xl/workbook.xml').decode('utf-8','ignore') if 'xl/workbook.xml' in names else ''
        defined=[]
        for attrs,val in re.findall(r'<definedName\b([^>]*)>(.*?)</definedName>',workbook,re.S):
            nm=re.search(r'name="([^"]+)"',attrs); hidden=re.search(r'hidden="([^"]+)"',attrs)
            defined.append({'name':nm.group(1) if nm else None,'hidden':hidden.group(1) if hidden else None,'value':re.sub(r'\s+',' ',val)[:300]})
        styles=z.read('xl/styles.xml').decode('utf-8','ignore') if 'xl/styles.xml' in names else ''
        numfmts=re.findall(r'<numFmt\b[^>]*numFmtId="([^"]+)"[^>]*formatCode="([^"]+)"',styles)
        return {'kind':'xlsx','core':core,'sheets':sheets,'defined_names':defined,'styles_hash':sha256_bytes(styles.encode()) if styles else None,
                'num_formats':numfmts,'formula_patterns':collections.Counter(formulas).most_common(100),'zip_times':zip_times(z),'text_style':text_style('\n'.join(text))}

def fingerprint(p:Path):
    base={'path':str(p),'name':p.name,'ext':p.suffix.lower(),'bytes':p.stat().st_size,'sha256':sha256_file(p)}
    try:
        if p.suffix.lower()=='.pdf': base.update(pdf_fp(p))
        elif p.suffix.lower()=='.docx': base.update(docx_fp(p))
        elif p.suffix.lower() in {'.xlsx','.xlsm'}: base.update(xlsx_fp(p))
    except Exception as e: base['error']=f'{type(e).__name__}: {e}'
    return base

def jacc(a,b):
    a=set(a or []); b=set(b or [])
    return len(a&b)/len(a|b) if a and b else 0.0

def build_prevalence(fps):
    stats={}
    for k in ('pdf','docx','xlsx'):
        group=[x for x in fps if x.get('kind')==k and not x.get('error')]; counts=collections.defaultdict(collections.Counter)
        for x in group:
            if k=='pdf':
                m=x.get('metadata',{})
                for key in ('creator','producer','author'):
                    if m.get(key): counts[key][m[key]]+=1
                if x.get('page_geometry'):counts['page_geometry'][json.dumps(x['page_geometry'],sort_keys=True)]+=1
                for v in set(tuple(z) for z in x.get('font_pairs',[])):counts['font_pairs'][v]+=1
                for v in set(x.get('font_subset_prefixes',[])):counts['font_subset_prefixes'][v]+=1
                for v,_ in x.get('jpeg_quant_hashes',[]):counts['jpeg_quant_hashes'][v]+=1
            elif k=='docx':
                c=x.get('core',{})
                for key in ('author','last_modified_by'):
                    if c.get(key):counts[key][c[key]]+=1
                for key in ('styles_hash','theme_hash','numbering_hash'):
                    if x.get(key):counts[key][x[key]]+=1
                for v in set(x.get('template_targets',[])):counts['template_targets'][v]+=1
                for v in set(x.get('rsids',[])):counts['rsids'][v]+=1
            else:
                c=x.get('core',{})
                for key in ('creator','lastModifiedBy'):
                    if c.get(key):counts[key][c[key]]+=1
                if x.get('styles_hash'):counts['styles_hash'][x['styles_hash']]+=1
                for d in x.get('defined_names',[]):counts['defined_names'][(d.get('name'),d.get('hidden'))]+=1
                for v in set(tuple(z) for z in x.get('num_formats',[])):counts['num_formats'][v]+=1
        stats[k]={'N':len(group),'counts':counts}
    return stats

def rarity(stats,kind,feature,value):
    if not stats or kind not in stats:return 1.0
    N=stats[kind]['N']
    if N<=1:return 1.0
    f=stats[kind]['counts'].get(feature,{}).get(value,1)
    return max(.08,math.log((N+1)/(f+0.5))/math.log(N+1))

def avg_intersection_rarity(stats,kind,feature,aa,bb):
    inter=set(aa or [])&set(bb or [])
    return sum(rarity(stats,kind,feature,v) for v in inter)/len(inter) if inter else 0.0

def score_pair(a,b,stats=None):
    score=0.0; reasons=[]
    if a['sha256']==b['sha256']:return 1.0,['byte-identical SHA-256']
    if a.get('kind')!=b.get('kind'):return 0.0,[]
    k=a.get('kind')
    if k=='pdf':
        ma=a.get('metadata',{}); mb=b.get('metadata',{})
        for key,w in [('creator',.10),('producer',.08),('author',.08)]:
            if ma.get(key) and ma.get(key)==mb.get(key):
                rr=rarity(stats,'pdf',key,ma[key]);score+=w*rr;reasons.append(f'same PDF {key}: {ma[key]} (rarity {rr:.2f})')
        if a.get('page_geometry')==b.get('page_geometry') and a.get('page_geometry'):
            rr=rarity(stats,'pdf','page_geometry',json.dumps(a['page_geometry'],sort_keys=True));score+=.08*rr;reasons.append(f'identical page geometry sequence (rarity {rr:.2f})')
        fj=jacc(a.get('font_pairs'),b.get('font_pairs'))
        if fj>.65:
            rr=avg_intersection_rarity(stats,'pdf','font_pairs',[tuple(z) for z in a.get('font_pairs',[])],[tuple(z) for z in b.get('font_pairs',[])]);score+=.10*fj*rr;reasons.append(f'font-set Jaccard {fj:.2f}, rarity {rr:.2f}')
        sj=jacc(a.get('font_subset_prefixes'),b.get('font_subset_prefixes'))
        if sj>.25:
            rr=avg_intersection_rarity(stats,'pdf','font_subset_prefixes',a.get('font_subset_prefixes'),b.get('font_subset_prefixes'));score+=.14*sj*rr;reasons.append(f'embedded-font subset-prefix Jaccard {sj:.2f}, rarity {rr:.2f}')
        qj=jacc([x[0] for x in a.get('jpeg_quant_hashes',[])],[x[0] for x in b.get('jpeg_quant_hashes',[])])
        if qj>.2:
            rr=avg_intersection_rarity(stats,'pdf','jpeg_quant_hashes',[x[0] for x in a.get('jpeg_quant_hashes',[])],[x[0] for x in b.get('jpeg_quant_hashes',[])]);score+=.14*qj*rr;reasons.append(f'JPEG quantization-table Jaccard {qj:.2f}, rarity {rr:.2f}')
    elif k=='docx':
        ca=a.get('core',{}); cb=b.get('core',{})
        for key,w in [('author',.10),('last_modified_by',.12)]:
            if ca.get(key) and ca.get(key)==cb.get(key):
                rr=rarity(stats,'docx',key,ca[key]);score+=w*rr;reasons.append(f'same DOCX {key}: {ca[key]} (rarity {rr:.2f})')
        for key,w in [('styles_hash',.10),('theme_hash',.05),('numbering_hash',.05)]:
            if a.get(key) and a.get(key)==b.get(key):
                rr=rarity(stats,'docx',key,a[key]);score+=w*rr;reasons.append(f'same {key} (rarity {rr:.2f})')
        rj=jacc(a.get('rsids'),b.get('rsids'))
        if rj>.2:
            rr=avg_intersection_rarity(stats,'docx','rsids',a.get('rsids'),b.get('rsids'));score+=.12*rj*rr;reasons.append(f'RSID Jaccard {rj:.2f}, rarity {rr:.2f}')
    elif k=='xlsx':
        ca=a.get('core',{}); cb=b.get('core',{})
        for key,w in [('creator',.10),('lastModifiedBy',.12)]:
            if ca.get(key) and ca.get(key)==cb.get(key):
                rr=rarity(stats,'xlsx',key,ca[key]);score+=w*rr;reasons.append(f'same XLSX {key}: {ca[key]} (rarity {rr:.2f})')
        if a.get('styles_hash') and a.get('styles_hash')==b.get('styles_hash'):
            rr=rarity(stats,'xlsx','styles_hash',a['styles_hash']);score+=.14*rr;reasons.append(f'identical XLSX styles.xml (rarity {rr:.2f})')
        da={(x['name'],x.get('hidden')) for x in a.get('defined_names',[])};db={(x['name'],x.get('hidden')) for x in b.get('defined_names',[])}
        dj=jacc(da,db)
        if dj>.2:
            rr=avg_intersection_rarity(stats,'xlsx','defined_names',da,db);score+=.16*dj*rr;reasons.append(f'defined-name Jaccard {dj:.2f}, rarity {rr:.2f}')
    ta=a.get('text_style',{});tb=b.get('text_style',{});sh=jacc(ta.get('shingle_sample'),tb.get('shingle_sample'))
    if sh>.08:score+=min(.18,.18*sh);reasons.append(f'rare 4/5-word shingle Jaccard {sh:.2f}')
    return min(1.0,score),reasons

def collect(args):
    out=[]
    for s in args:
        p=Path(s)
        if p.is_file() and p.suffix.lower() in SUPPORTED:out.append(p)
        elif p.is_dir():out.extend(q for q in p.rglob('*') if q.is_file() and q.suffix.lower() in SUPPORTED)
    return sorted(set(out))

def main():
    ap=argparse.ArgumentParser();ap.add_argument('paths',nargs='+');ap.add_argument('--out',default='fingerprints.json');ap.add_argument('--csv',default=None);ap.add_argument('--pairs',type=float,default=.18)
    ns=ap.parse_args();fps=[fingerprint(p) for p in collect(ns.paths)];stats=build_prevalence(fps);pairs=[]
    for i in range(len(fps)):
        for j in range(i+1,len(fps)):
            s,r=score_pair(fps[i],fps[j],stats)
            if s>=ns.pairs:pairs.append({'a':fps[i]['path'],'b':fps[j]['path'],'score':round(s,4),'reasons':r})
    pairs.sort(key=lambda x:-x['score'])
    result={'warning':'Similarity is a screening lead, not evidence of collusion/common control. Falsify shared templates, common submission portals, scanners, consultants, legitimate partnerships and agency-side recombination first.','files':fps,'candidate_pairs':pairs}
    Path(ns.out).write_text(json.dumps(result,indent=2,ensure_ascii=False,default=str))
    if ns.csv:
        import csv
        with open(ns.csv,'w',newline='',encoding='utf-8') as f:
            w=csv.writer(f);w.writerow(['score','file_a','file_b','reasons'])
            for x in pairs:w.writerow([x['score'],x['a'],x['b'],' | '.join(x['reasons'])])
    print(f'fingerprinted {len(fps)} files; {len(pairs)} candidate pairs >= {ns.pairs:.2f}')
    for x in pairs[:20]:print(f"{x['score']:.3f}  {Path(x['a']).name} <> {Path(x['b']).name}\n       "+'; '.join(x['reasons'][:6]))

if __name__=='__main__':main()
